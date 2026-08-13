"""
crawl_routes.py — FastAPI router for SentinelAI Prompt Crawler endpoints.

Endpoints:
  POST /api/v1/crawl/prompt              — Multi-platform prompt crawl
  POST /api/v1/crawl/single/{platform}   — Single-platform test
  POST /api/v1/crawl/account             — Target user handle crawl
  POST /api/v1/crawl/watchlist/add       — Add to continuous watchlist
  GET  /api/v1/crawl/status              — Crawler health & stats
  POST /api/v1/crawl/export              — Export crawl results as PDF or Excel
"""
from __future__ import annotations

import datetime
import io
import os
import re
from typing import Any, Dict, List, Optional

import asyncio
import json
from fastapi import APIRouter, BackgroundTasks, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field

import re

ILLEGAL_CHARACTERS_RE = re.compile(r'[\000-\010]|[\013-\014]|[\016-\037]')

def sanitize(val):
    if isinstance(val, str):
        return ILLEGAL_CHARACTERS_RE.sub('', val)
    return val

from crawlers.hybrid_crawler import get_hybrid_crawler, ALL_PLATFORMS
from nlp_service.models.inference import run_nlp_pipeline
from storage.db_client import db_client

router = APIRouter(prefix="/crawl", tags=["Prompt Crawler"])


# ── Request / Response models ─────────────────────────────────────────────────

class PromptCrawlRequest(BaseModel):
    prompt: str = Field(..., min_length=1, max_length=500, example="india protest news")
    platforms: Optional[List[str]] = Field(
        default=None,
        example=["X", "YouTube", "Telegram", "Instagram"],
        description="Platforms to crawl. Defaults to all available if omitted.",
    )
    limit: int = Field(default=20, ge=1, le=5000, description="Max results per platform")
    time_filter: str = Field(default="any", description="Time filter: any, 24h, 48h, 1week, 1month")
    fetch_comments: bool = Field(
        default=False,
        description="Fetch top comments from YouTube/X",
    )


class SinglePlatformRequest(BaseModel):
    prompt: str = Field(..., min_length=1, max_length=500)
    limit: int = Field(default=20, ge=1, le=5000)
    time_filter: str = "any"
    fetch_comments: bool = False


class AccountCrawlRequest(BaseModel):
    account_handle: str = Field(..., min_length=1, max_length=100, example="@target_user")
    platform: Optional[str] = Field(default="all", example="x")
    limit: int = Field(default=20, ge=1, le=5000)
    time_filter: str = "any"


class WatchlistAddRequest(BaseModel):
    query: str = Field(..., min_length=1, max_length=500, example="covid updates")
    platforms: List[str] = Field(default=["X", "YouTube"])
    frequency_minutes: int = Field(default=15, ge=1, le=1440)
    priority: str = Field(default="normal", pattern="^(high|normal|low)$")


# ── Helper for processing harvested posts via buildspec NLP ───────────────────

def _enrich_and_store_crawled_posts(posts: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    enriched = []
    for post in posts:
        post_id = str(post.get("id") or post.get("post_id") or f"crawled_{datetime.datetime.utcnow().timestamp()}")
        content = post.get("content") or post.get("text") or ""
        
        if content:
            nlp_res = run_nlp_pipeline(post_id=post_id, text=content)
            post["nlp_analysis"] = nlp_res
            post["threat_level"] = nlp_res.get("threat_category", {}).get("label", "Neutral")
            post["sentiment"] = nlp_res.get("sentiment", {}).get("label", "neutral")
            post["is_hate_speech"] = nlp_res.get("hate_speech", {}).get("flag", False)
            post["requires_review"] = nlp_res.get("requires_human_review", False)
        
        db_client.save_post(post)
        enriched.append(post)
    return enriched

# ── Endpoints ────────────────────────────────────────────────────────────────────

_EXPORTS_DIR = os.path.join(os.path.dirname(__file__), "..", "..", "data", "exports")


def _auto_save_excel(result: Dict[str, Any]) -> None:
    """
    Automatically save crawl results to an Excel file in data/exports/.
    Called as a background task — never blocks the API response.
    Filename pattern: crawl_<query>_<timestamp>.xlsx
    """
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        os.makedirs(_EXPORTS_DIR, exist_ok=True)

        query      = result.get("query", "unknown")
        timestamp  = datetime.datetime.utcnow().strftime("%Y%m%d_%H%M%S")
        safe_query = re.sub(r"[^\w\s-]", "", query).strip().replace(" ", "_")[:40]
        filename   = f"crawl_{safe_query}_{timestamp}.xlsx"
        filepath   = os.path.join(_EXPORTS_DIR, filename)

        wb = openpyxl.Workbook()

        # ── Style palette
        HDR_FILL  = PatternFill("solid", fgColor="0F172A")
        ROW_FILL1 = PatternFill("solid", fgColor="0F172A")
        ROW_FILL2 = PatternFill("solid", fgColor="111827")
        HDR_FONT  = Font(bold=True, color="60A5FA", size=10)
        CELL_FONT = Font(color="E2E8F0", size=9)
        TITLE_FONT = Font(bold=True, color="3B82F6", size=14)
        META_FONT  = Font(color="E2E8F0", size=10)
        WRAP       = Alignment(wrap_text=True, vertical="top")
        thin       = Side(style="thin", color="1E3A5F")
        border     = Border(left=thin, right=thin, top=thin, bottom=thin)

        COLUMNS = [
            ("#",            5),
            ("Platform",     14),
            ("Author",       20),
            ("Content",      65),
            ("URL",          45),
            ("Language",     10),
            ("Threat Level", 15),
            ("Sentiment",    13),
            ("Likes",        10),
            ("Shares",       10),
            ("Views",        10),
            ("Comments",     10),
            ("Timestamp",    22),
        ]

        platform_results = result.get("results", {})

        # ── Summary sheet
        ws_sum = wb.active
        ws_sum.title = "Summary"
        ws_sum.sheet_view.showGridLines = False
        ws_sum["A1"] = "SentinelAI Crawl Report"
        ws_sum["A1"].font = TITLE_FONT
        ws_sum["A2"] = f"Query: {query}"
        ws_sum["A2"].font = META_FONT
        ws_sum["A3"] = f"Timestamp: {result.get('timestamp', timestamp)}"
        ws_sum["A3"].font = META_FONT
        ws_sum["A4"] = (
            f"Total Posts: {result.get('total_posts', 0)}  |  "
            f"Platforms: {result.get('platforms_crawled', 0)}  |  "
            f"Duration: {result.get('duration_s', 0)}s"
        )
        ws_sum["A4"].font = META_FONT
        ws_sum.column_dimensions["A"].width = 80

        ov_headers = ["Platform", "Status", "Post Count", "Duration (s)"]
        for ci, h in enumerate(ov_headers, 1):
            cell = ws_sum.cell(row=6, column=ci, value=h)
            cell.font = HDR_FONT
            cell.fill = HDR_FILL
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for ri, (plat, pdata) in enumerate(platform_results.items(), 7):
            fill = ROW_FILL1 if ri % 2 == 0 else ROW_FILL2
            for ci, val in enumerate(
                [plat, pdata.get("status", "N/A"), pdata.get("count", 0), pdata.get("duration_s", 0)], 1
            ):
                cell = ws_sum.cell(row=ri, column=ci, value=val)
                cell.font = CELL_FONT
                cell.fill = fill
                cell.border = border
                cell.alignment = Alignment(horizontal="left", vertical="top")
        for ci in range(1, 5):
            ws_sum.column_dimensions[get_column_letter(ci)].width = 22

        # ── All Posts sheet
        ws_all = wb.create_sheet(title="All Posts")
        ws_all.sheet_view.showGridLines = False
        for ci, (col_name, col_width) in enumerate(COLUMNS, 1):
            cell = ws_all.cell(row=1, column=ci, value=col_name)
            cell.font = HDR_FONT
            cell.fill = HDR_FILL
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center")
            ws_all.column_dimensions[get_column_letter(ci)].width = col_width
        ws_all.row_dimensions[1].height = 20

        row_idx = 2
        for plat, pdata in platform_results.items():
            for post in pdata.get("posts", []):
                eng  = post.get("engagement") or {}
                fill = ROW_FILL1 if row_idx % 2 == 0 else ROW_FILL2
                vals = [
                    row_idx - 1,
                    plat,
                    post.get("author_username") or post.get("author") or "",
                    (post.get("content") or post.get("text") or "")[:2000],
                    post.get("url") or "",
                    (post.get("language") or "").upper(),
                    post.get("threat_level") or "Neutral",
                    post.get("sentiment") or "neutral",
                    eng.get("likes") or 0,
                    eng.get("shares") or 0,
                    eng.get("views") or 0,
                    eng.get("comments") or 0,
                    post.get("timestamp") or post.get("crawled_at") or "",
                ]
                for ci, val in enumerate(vals, 1):
                    cell = ws_all.cell(row=row_idx, column=ci, value=sanitize(val))
                    cell.font = CELL_FONT
                    cell.fill = fill
                    cell.border = border
                    cell.alignment = WRAP
                ws_all.row_dimensions[row_idx].height = 50
                row_idx += 1

        # ── Per-platform sheets
        for plat, pdata in platform_results.items():
            ws = wb.create_sheet(title=plat[:31])
            ws.sheet_view.showGridLines = False
            for ci, (col_name, col_width) in enumerate(COLUMNS, 1):
                cell = ws.cell(row=1, column=ci, value=col_name)
                cell.font = HDR_FONT
                cell.fill = HDR_FILL
                cell.border = border
                cell.alignment = Alignment(horizontal="center", vertical="center")
                ws.column_dimensions[get_column_letter(ci)].width = col_width
            ws.row_dimensions[1].height = 20

            for rn, post in enumerate(pdata.get("posts", []), 2):
                eng  = post.get("engagement") or {}
                fill = ROW_FILL1 if rn % 2 == 0 else ROW_FILL2
                vals = [
                    rn - 1,
                    plat,
                    post.get("author_username") or post.get("author") or "",
                    (post.get("content") or post.get("text") or "")[:2000],
                    post.get("url") or "",
                    (post.get("language") or "").upper(),
                    post.get("threat_level") or "Neutral",
                    post.get("sentiment") or "neutral",
                    eng.get("likes") or 0,
                    eng.get("shares") or 0,
                    eng.get("views") or 0,
                    eng.get("comments") or 0,
                    post.get("timestamp") or post.get("crawled_at") or "",
                ]
                for ci, val in enumerate(vals, 1):
                    cell = ws.cell(row=rn, column=ci, value=sanitize(val))
                    cell.font = CELL_FONT
                    cell.fill = fill
                    cell.border = border
                    cell.alignment = WRAP
                ws.row_dimensions[rn].height = 50

        # Make "All Posts" the active sheet when opening
        if "All Posts" in wb.sheetnames:
            wb.active = wb.sheetnames.index("All Posts")

        wb.save(filepath)
        total = sum(len(v.get("posts", [])) for v in platform_results.values())
        print(f"[AutoExport] ✅ Saved {total} posts → {filepath}")

    except Exception as exc:
        import traceback
        print(f"[AutoExport] ⚠️ Failed to auto-save Excel: {exc}")
        traceback.print_exc()


@router.post("/prompt", summary="Multi-platform prompt crawl")
async def crawl_by_prompt(
    request: PromptCrawlRequest,
    background_tasks: BackgroundTasks,
) -> Dict[str, Any]:
    """
    Crawl selected platforms in parallel with query prompt
    and process all harvested posts through buildspec NLP threat/sentiment pipeline.
    """
    platforms = request.platforms or ALL_PLATFORMS
    crawler = get_hybrid_crawler()

    result = await crawler.crawl_multi_platform(
        queries=[request.prompt],
        platforms=platforms,
        limit=request.limit,
        fetch_comments=request.fetch_comments,
        time_filter=request.time_filter,
    )

    if "error" in result and not result.get("results"):
        raise HTTPException(status_code=400, detail=result["error"])

    # Enrich all harvested posts with real NLP inference & store centrally
    for plat, plat_res in result.get("results", {}).items():
        if isinstance(plat_res, dict) and "posts" in plat_res:
            raw_posts = plat_res.get("posts", [])
            if raw_posts:
                enriched = await asyncio.to_thread(_enrich_and_store_crawled_posts, raw_posts)
                result["results"][plat]["posts"] = enriched

    # Auto-save results to Excel in background (non-blocking)
    background_tasks.add_task(_auto_save_excel, result)

    return result


@router.post("/prompt/stream", summary="Stream Multi-platform prompt crawl")
async def stream_crawl_by_prompt(
    request: PromptCrawlRequest,
    background_tasks: BackgroundTasks,
):
    """
    Stream platform results as they complete using NDJSON.
    """
    platforms = request.platforms or ALL_PLATFORMS
    crawler = get_hybrid_crawler()

    async def generate():
        import time
        global_start = time.monotonic()
        total_posts = 0
        platforms_crawled = 0
        accumulated_results: Dict[str, Any] = {}
        
        async for platform, plat_res in crawler.stream_multi_platform(
            queries=[request.prompt],
            platforms=platforms,
            limit=request.limit,
            fetch_comments=request.fetch_comments,
            time_filter=request.time_filter,
        ):
            platforms_crawled += 1
            if isinstance(plat_res, dict) and "posts" in plat_res:
                raw_posts = plat_res.get("posts", [])
                if raw_posts:
                    enriched = await asyncio.to_thread(_enrich_and_store_crawled_posts, raw_posts)
                    plat_res["posts"] = enriched
                total_posts += plat_res.get("count", 0)
            
            accumulated_results[platform] = plat_res
            yield json.dumps({"type": "platform_result", "platform": platform, "data": plat_res}) + "\n"
        
        total_elapsed = round(time.monotonic() - global_start, 2)
        summary_payload = {
            "type": "summary",
            "duration_s": total_elapsed,
            "total_posts": total_posts,
            "platforms_crawled": platforms_crawled
        }
        yield json.dumps(summary_payload) + "\n"

        # Auto-save full result to Excel once the stream is complete
        full_result = {
            "query": request.prompt,
            "timestamp": datetime.datetime.utcnow().isoformat(),
            "total_posts": total_posts,
            "platforms_crawled": platforms_crawled,
            "duration_s": total_elapsed,
            "results": accumulated_results,
        }
        await asyncio.to_thread(_auto_save_excel, full_result)

    return StreamingResponse(generate(), media_type="application/x-ndjson")


@router.post("/single/{platform}", summary="Single-platform test crawl")
async def crawl_single_platform(
    platform: str,
    request: SinglePlatformRequest,
    background_tasks: BackgroundTasks,
) -> Dict[str, Any]:
    """
    Crawl a single platform for targeted queries and process through buildspec NLP pipeline.
    """
    if platform not in ALL_PLATFORMS:
        raise HTTPException(
            status_code=400,
            detail=f"Unknown platform '{platform}'. Valid: {ALL_PLATFORMS}",
        )

    crawler = get_hybrid_crawler()
    result = await crawler.crawl_platform(
        query=request.prompt,
        platform=platform,
        limit=request.limit,
        fetch_comments=request.fetch_comments,
        time_filter=request.time_filter,
    )

    if result.get("posts"):
        result["posts"] = await asyncio.to_thread(_enrich_and_store_crawled_posts, result["posts"])

    return {
        "query": request.prompt,
        "platforms_crawled": 1,
        "total_posts": result.get("count", 0),
        "duration_s": result.get("duration_s", 0),
        "timestamp": datetime.datetime.utcnow().isoformat(),
        "results": {platform: result},
    }


@router.post("/account", summary="Targeted account handle/ID crawl")
async def crawl_by_account(request: AccountCrawlRequest) -> Dict[str, Any]:
    """
    Targeted Account Endpoint — crawl user posts from X, Instagram, Facebook, Telegram, or YouTube.
    """
    crawler = get_hybrid_crawler()
    target_platform = request.platform if request.platform and request.platform != "all" else "X"
    
    result = await crawler.crawl_platform(
        query=request.account_handle,
        platform=target_platform,
        limit=request.limit,
        time_filter=request.time_filter,
    )

    if result.get("posts"):
        result["posts"] = await asyncio.to_thread(_enrich_and_store_crawled_posts, result["posts"])

    return {
        "account": request.account_handle,
        "platform": target_platform,
        "count": result.get("count", 0),
        "posts": result.get("posts", []),
        "status": result.get("status", "success"),
    }


@router.post("/watchlist/add", summary="Add query to continuous watchlist")
async def add_to_watchlist(request: WatchlistAddRequest) -> Dict[str, Any]:
    """
    Register a query for continuous background monitoring.
    """
    invalid = [p for p in request.platforms if p not in ALL_PLATFORMS]
    if invalid:
        raise HTTPException(
            status_code=400,
            detail=f"Invalid platforms: {invalid}. Valid: {ALL_PLATFORMS}",
        )

    crawler = get_hybrid_crawler()
    return crawler.add_to_watchlist(
        query=request.query,
        platforms=request.platforms,
        frequency_minutes=request.frequency_minutes,
        priority=request.priority,
    )


@router.get("/status", summary="Crawler health & statistics")
async def get_crawler_status() -> Dict[str, Any]:
    """
    Returns real-time health information for the HybridCrawler.
    """
    crawler = get_hybrid_crawler()
    return crawler.get_status()


# ── Export models ──────────────────────────────────────────────────────────────

class ExportRequest(BaseModel):
    format: str = Field(default="pdf", pattern="^(pdf|excel)$", description="Export format: 'pdf' or 'excel'")
    results: Dict[str, Any] = Field(..., description="Full crawl results payload from /crawl/prompt")


# ── PDF generator ──────────────────────────────────────────────────────────────

def _generate_pdf(results: Dict[str, Any]) -> bytes:
    """Generate a styled PDF report from crawl results using reportlab."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,
        HRFlowable, KeepTogether
    )
    from reportlab.lib.enums import TA_LEFT, TA_CENTER

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        rightMargin=1.8 * cm,
        leftMargin=1.8 * cm,
        topMargin=2 * cm,
        bottomMargin=2 * cm,
        title="SentinelAI Crawl Report",
        author="SentinelAI",
    )

    styles = getSampleStyleSheet()
    # Custom styles
    title_style = ParagraphStyle(
        "ReportTitle",
        parent=styles["Title"],
        fontSize=22,
        textColor=colors.HexColor("#3b82f6"),
        spaceAfter=6,
        alignment=TA_CENTER,
    )
    subtitle_style = ParagraphStyle(
        "Subtitle",
        parent=styles["Normal"],
        fontSize=10,
        textColor=colors.HexColor("#94a3b8"),
        spaceAfter=4,
        alignment=TA_CENTER,
    )
    section_style = ParagraphStyle(
        "SectionHeader",
        parent=styles["Heading2"],
        fontSize=13,
        textColor=colors.HexColor("#60a5fa"),
        spaceBefore=14,
        spaceAfter=6,
        borderPad=4,
    )
    post_content_style = ParagraphStyle(
        "PostContent",
        parent=styles["Normal"],
        fontSize=8.5,
        textColor=colors.HexColor("#e2e8f0"),
        leading=13,
        spaceAfter=2,
    )
    meta_style = ParagraphStyle(
        "MetaInfo",
        parent=styles["Normal"],
        fontSize=7.5,
        textColor=colors.HexColor("#64748b"),
        spaceAfter=1,
    )
    threat_colors = {
        "Critical": colors.HexColor("#ef4444"),
        "High": colors.HexColor("#f97316"),
        "Medium": colors.HexColor("#eab308"),
        "Neutral": colors.HexColor("#22c55e"),
        "Low": colors.HexColor("#22c55e"),
    }

    query = results.get("query", "N/A")
    timestamp = results.get("timestamp", datetime.datetime.utcnow().isoformat())
    total_posts = results.get("total_posts", 0)
    platforms_crawled = results.get("platforms_crawled", 0)
    duration = results.get("duration_s", 0)
    platform_results = results.get("results", {})

    story = []

    # -- Cover section
    story.append(Spacer(1, 0.5 * cm))
    story.append(Paragraph("[SentinelAI] Crawl Intelligence Report", title_style))
    story.append(Paragraph("Powered by Anvaya Crawler &amp; NLP Threat Pipeline", subtitle_style))
    story.append(Spacer(1, 0.4 * cm))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#1e293b")))
    story.append(Spacer(1, 0.4 * cm))

    # Summary table
    summary_data = [
        ["Query", query],
        ["Timestamp", timestamp],
        ["Total Posts", str(total_posts)],
        ["Platforms Crawled", str(platforms_crawled)],
        ["Duration", f"{duration}s"],
    ]
    summary_table = Table(summary_data, colWidths=[4 * cm, 13 * cm])
    summary_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#0f172a")),
        ("BACKGROUND", (1, 0), (1, -1), colors.HexColor("#1e293b")),
        ("TEXTCOLOR", (0, 0), (0, -1), colors.HexColor("#60a5fa")),
        ("TEXTCOLOR", (1, 0), (1, -1), colors.HexColor("#e2e8f0")),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
        ("ROWBACKGROUNDS", (0, 0), (-1, -1), [colors.HexColor("#0f172a"), colors.HexColor("#1e293b")]),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#334155")),
        ("PADDING", (0, 0), (-1, -1), 6),
    ]))
    story.append(summary_table)
    story.append(Spacer(1, 0.6 * cm))

    # ── Per-platform sections
    for platform, plat_data in platform_results.items():
        posts = plat_data.get("posts", [])
        count = plat_data.get("count", 0)
        status = plat_data.get("status", "unknown")

        section_header = Paragraph(
            f"[{platform}]  {count} posts  [{status}]",
            section_style
        )
        story.append(section_header)
        story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#1e3a5f")))
        story.append(Spacer(1, 0.2 * cm))

        if not posts:
            story.append(Paragraph("No posts found for this platform.", meta_style))
            story.append(Spacer(1, 0.3 * cm))
            continue

        for i, post in enumerate(posts, 1):
            content = (post.get("content") or post.get("text") or "").strip()
            author = post.get("author_username") or post.get("author") or "unknown"
            url = post.get("url") or ""
            language = (post.get("language") or "N/A").upper()
            threat = post.get("threat_level") or "Neutral"
            sentiment = post.get("sentiment") or "neutral"
            eng = post.get("engagement") or {}
            likes = eng.get("likes", 0) or 0
            shares = eng.get("shares", 0) or 0
            views = eng.get("views", 0) or 0

            threat_color = threat_colors.get(threat, colors.HexColor("#94a3b8"))

            row_items = [
                Paragraph(f"<b>#{i} @{author}</b>", ParagraphStyle(
                    "PostAuthor", parent=styles["Normal"],
                    fontSize=8.5, textColor=colors.HexColor("#a78bfa"), spaceAfter=2
                )),
                Paragraph(content[:500] + ("..." if len(content) > 500 else ""), post_content_style),
                Paragraph(
                    f"Lang: {language} | Threat: {threat} | Sentiment: {sentiment} | "
                    f"Likes: {likes:,}  Shares: {shares:,}  Views: {views:,}",
                    meta_style
                ),
            ]
            if url:
                safe_url = url[:80] + ("..." if len(url) > 80 else "")
                row_items.append(Paragraph(
                    f"<link href='{url}'>{safe_url}</link>",
                    meta_style
                ))

            block = KeepTogether(row_items + [Spacer(1, 0.15 * cm)])
            story.append(block)

        story.append(Spacer(1, 0.3 * cm))

    # Footer note
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#1e293b")))
    story.append(Spacer(1, 0.3 * cm))
    story.append(Paragraph(
        f"Generated by SentinelAI on {datetime.datetime.utcnow().strftime('%Y-%m-%d %H:%M:%S')} UTC",
        subtitle_style
    ))

    doc.build(story)
    buf.seek(0)
    return buf.read()


# ── Excel generator ────────────────────────────────────────────────────────────

def _generate_excel(results: Dict[str, Any]) -> bytes:
    """Generate an Excel workbook from crawl results using openpyxl."""
    import openpyxl
    from openpyxl.styles import (
        Font, PatternFill, Alignment, Border, Side
    )
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()

    # ── Color palette
    HDR_FILL  = PatternFill("solid", fgColor="0F172A")
    META_FILL = PatternFill("solid", fgColor="1E293B")
    ROW_FILL1 = PatternFill("solid", fgColor="0F172A")
    ROW_FILL2 = PatternFill("solid", fgColor="111827")
    HDR_FONT  = Font(bold=True, color="60A5FA", size=10)
    META_FONT = Font(bold=False, color="94A3B8", size=9)
    CELL_FONT = Font(color="E2E8F0", size=9)
    WRAP      = Alignment(wrap_text=True, vertical="top")

    thin = Side(style="thin", color="1E3A5F")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    COLUMNS = [
        ("#", 5),
        ("Platform", 12),
        ("Author", 18),
        ("Content", 60),
        ("URL", 40),
        ("Language", 10),
        ("Threat Level", 14),
        ("Sentiment", 12),
        ("Likes", 10),
        ("Shares", 10),
        ("Views", 10),
        ("Comments", 10),
        ("Timestamp", 22),
    ]

    platform_results = results.get("results", {})
    query = results.get("query", "N/A")

    # ── Summary sheet
    ws_summary = wb.active
    ws_summary.title = "Summary"
    ws_summary.sheet_view.showGridLines = False
    ws_summary["A1"] = "SentinelAI Crawl Report"
    ws_summary["A1"].font = Font(bold=True, color="3B82F6", size=16)
    ws_summary["A2"] = f"Query: {query}"
    ws_summary["A2"].font = Meta_font = Font(color="E2E8F0", size=11)
    ws_summary["A3"] = f"Timestamp: {results.get('timestamp', 'N/A')}"
    ws_summary["A3"].font = Meta_font
    ws_summary["A4"] = f"Total Posts: {results.get('total_posts', 0)}  |  Platforms: {results.get('platforms_crawled', 0)}  |  Duration: {results.get('duration_s', 0)}s"
    ws_summary["A4"].font = Meta_font
    ws_summary.column_dimensions["A"].width = 80
    ws_summary["A1"].fill = PatternFill("solid", fgColor="0F172A")

    # Platform overview table
    overview_headers = ["Platform", "Status", "Post Count", "Duration (s)"]
    for col_idx, header in enumerate(overview_headers, 1):
        cell = ws_summary.cell(row=6, column=col_idx, value=header)
        cell.font = HDR_FONT
        cell.fill = HDR_FILL
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row_idx, (platform, plat_data) in enumerate(platform_results.items(), 7):
        fill = ROW_FILL1 if row_idx % 2 == 0 else ROW_FILL2
        for col_idx, val in enumerate([
            platform,
            plat_data.get("status", "N/A"),
            plat_data.get("count", 0),
            plat_data.get("duration_s", 0),
        ], 1):
            cell = ws_summary.cell(row=row_idx, column=col_idx, value=val)
            cell.font = CELL_FONT
            cell.fill = fill
            cell.border = border
            cell.alignment = Alignment(horizontal="left", vertical="top")
    for col_idx in range(1, 5):
        ws_summary.column_dimensions[get_column_letter(col_idx)].width = 20

    # ── All Posts sheet (flat view)
    ws_all = wb.create_sheet(title="All Posts")
    ws_all.sheet_view.showGridLines = False

    for col_idx, (col_name, col_width) in enumerate(COLUMNS, 1):
        cell = ws_all.cell(row=1, column=col_idx, value=col_name)
        cell.font = HDR_FONT
        cell.fill = HDR_FILL
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws_all.column_dimensions[get_column_letter(col_idx)].width = col_width
    ws_all.row_dimensions[1].height = 20

    row_idx = 2
    for platform, plat_data in platform_results.items():
        posts = plat_data.get("posts", [])
        for post in posts:
            eng = post.get("engagement") or {}
            fill = ROW_FILL1 if row_idx % 2 == 0 else ROW_FILL2
            values = [
                row_idx - 1,
                platform,
                post.get("author_username") or post.get("author") or "",
                (post.get("content") or post.get("text") or "")[:2000],
                post.get("url") or "",
                (post.get("language") or "").upper(),
                post.get("threat_level") or "Neutral",
                post.get("sentiment") or "neutral",
                eng.get("likes") or 0,
                eng.get("shares") or 0,
                eng.get("views") or 0,
                eng.get("comments") or 0,
                post.get("timestamp") or "",
            ]
            for col_idx, val in enumerate(values, 1):
                cell = ws_all.cell(row=row_idx, column=col_idx, value=sanitize(val))
                cell.font = CELL_FONT
                cell.fill = fill
                cell.border = border
                cell.alignment = WRAP
            ws_all.row_dimensions[row_idx].height = 48
            row_idx += 1

    # ── Per-platform sheets
    for platform, plat_data in platform_results.items():
        sheet_name = platform[:31]  # Excel sheet name limit
        ws = wb.create_sheet(title=sheet_name)
        ws.sheet_view.showGridLines = False

        for col_idx, (col_name, col_width) in enumerate(COLUMNS, 1):
            cell = ws.cell(row=1, column=col_idx, value=col_name)
            cell.font = HDR_FONT
            cell.fill = HDR_FILL
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center")
            ws.column_dimensions[get_column_letter(col_idx)].width = col_width
        ws.row_dimensions[1].height = 20

        posts = plat_data.get("posts", [])
        for row_num, post in enumerate(posts, 2):
            eng = post.get("engagement") or {}
            fill = ROW_FILL1 if row_num % 2 == 0 else ROW_FILL2
            values = [
                row_num - 1,
                platform,
                post.get("author_username") or post.get("author") or "",
                (post.get("content") or post.get("text") or "")[:2000],
                post.get("url") or "",
                (post.get("language") or "").upper(),
                post.get("threat_level") or "Neutral",
                post.get("sentiment") or "neutral",
                eng.get("likes") or 0,
                eng.get("shares") or 0,
                eng.get("views") or 0,
                eng.get("comments") or 0,
                post.get("timestamp") or "",
            ]
            for col_idx, val in enumerate(values, 1):
                cell = ws.cell(row=row_num, column=col_idx, value=sanitize(val))
                cell.font = CELL_FONT
                cell.fill = fill
                cell.border = border
                cell.alignment = WRAP
            ws.row_dimensions[row_num].height = 48

    # Make "All Posts" the active sheet when opening
    wb.active = wb.sheetnames.index("All Posts")

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


# ── Export endpoint ────────────────────────────────────────────────────────────

@router.post("/export", summary="Export crawl results as PDF or Excel")
async def export_crawl_results(request: ExportRequest):
    """
    Accept a full crawl result payload and return it as a downloadable
    PDF report or Excel workbook.

    - **format=pdf**   → styled PDF via reportlab
    - **format=excel** → .xlsx workbook via openpyxl
    """
    fmt = request.format.lower()

    if fmt == "pdf":
        pdf_bytes = await asyncio.to_thread(_generate_pdf, request.results)
        filename = f"sentinelai_crawl_{datetime.datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.pdf"
        return StreamingResponse(
            io.BytesIO(pdf_bytes),
            media_type="application/pdf",
            headers={"Content-Disposition": f'attachment; filename="{filename}"'},
        )

    elif fmt == "excel":
        xlsx_bytes = await asyncio.to_thread(_generate_excel, request.results)
        filename = f"sentinelai_crawl_{datetime.datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.xlsx"
        return StreamingResponse(
            io.BytesIO(xlsx_bytes),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{filename}"'},
        )

    raise HTTPException(status_code=400, detail="Invalid format. Use 'pdf' or 'excel'.")
